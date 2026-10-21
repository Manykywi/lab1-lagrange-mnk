from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "AI Interaction Log"
F = "Arial"
thin = Side(style="thin", color="000000")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
green = PatternFill("solid", fgColor="00FF00")

headers = ["№п.п", "Компонент системи", "Назва задачі", "Хто розробляв (Самостійно / AI / Спільно)", "Коментар / Промт"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(name=F, bold=True, size=11)
    cell.fill = green
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

P1 = 'Промт 1: «Розробити програмний продукт відповідно до варіанту завдання лабораторної роботи № 1 …» (див. docs/ai-log.md, 8.2)'
rows = [
 ("Завдання", "Отримання варіанта та вимог (варіант 6, набір із 5 точок)", "Самостійно", "Варіант, методичку та скріншоти (таблиця варіантів, зразок AI-логу) надав студент."),
 ("ТЗ", "Опис технічного завдання: FR/NFR, інтерфейс, обмеження, критерії приймання", "Спільно", P1 + "\nКоментар: вимоги — з методички студента; формулювання ТЗ та припущення A1 (набори 10/20 точок) — AI."),
 ("Вибір методології", "Agile: 5 ітерацій, порівняння з Waterfall / V-model / Spiral", "AI", "Запропоновано AI, обґрунтування в docs/02-methodology.md; студент підтверджує вибір."),
 ("Проєктування логіки", "Блок-схеми: головний сценарій, Лагранж, МНК (QR), автомасштаб, цикл анімації", "AI", "Mermaid → PNG (tools/build-diagrams.js)."),
 ("Проєктування логіки", "Таблиці ідентифікаторів: змінні, функції, об'єкти", "AI", "docs/03-design.md, п. 3.7."),
 ("Проєктування логіки", "Алгоритм Лагранжа (+ форма Ньютона, поділені різниці)", "AI", "Промт 1; src/math.js: lagrangeEval, dividedDifferences, newtonEval."),
 ("Програмна реалізація", "МНК: дизайн-матриця, QR-розклад Хаусхолдера, нормування x, залишки", "AI", "Промт 4: «не використовувати нормальні рівняння у підсумкових розрахунках»; fitPolynomial, solveLeastSquares."),
 ("Програмна реалізація", "Автомасштабування осей X та Y за min/max функції", "AI", "Промт 1 (примітка про автомасштаб). Chart.computeTarget, LabMath.autoScale; працює і під час анімації."),
 ("Програмна реалізація", "Canvas-графіки: осі, сітка, легенда, підказки, адаптивність, теми", "AI", "src/chart.js, src/style.css."),
 ("Програмна реалізація", "Анімація інтерполяції та МНК", "AI", "Промт 5. src/animation.js + src/app.js: покрокове додавання вузлів, ріст полінома від ȳ, відрізки залишків."),
 ("Програмна реалізація", "6 графіків (3 Лагранж + 3 МНК) та перемикач наборів 5/10/20", "AI", "Подання «Порівняння 5·10·20» та кнопки наборів."),
 ("Тестування", "Автотести ядра, даних і таймлайнів (node:test)", "AI", "tests/*.test.js. Критичний аналіз: тест-«тавтологія» виправлено (docs/ai-log.md, випадок 1)."),
 ("Тестування", "Порівняльний аналіз для 5/10/20 точок (похибки, LOO, Рунге, QR vs нормальні рівняння)", "AI", "tools/analysis.js → docs/05-analysis.md. Знайдено й виправлено 2 помилки (випадки 2 і 3)."),
 ("Моделювання взаємодії", "Діаграми послідовності та компонентна діаграма", "AI", "docs/04-interaction.md, docs/03-design.md."),
 ("Документування", "Звіт (.docx), README, AI-лог", "AI", "Текст і оформлення підготував AI за вимогами завдання."),
 ("Документування", "Локальна історія git-комітів по ітераціях методології", "AI", "Коміти створено AI; публікація на GitHub — студент."),
 ("Контроль якості", "Запуск застосунку, перегляд графіків, перевірка результатів і висновків", "Самостійно", "Виконує студент перед здачею: index.html, npm test, звірка висновків зі своїм розумінням теми."),
 ("Публікація", "Репозиторій GitHub, публічне посилання на чат, скріншоти діалогу", "Самостійно", "Виконує студент (інструкція в README.md та docs/ai-log.md, 8.3)."),
]
for i, (comp, task, who, comm) in enumerate(rows, 1):
    r = i + 1
    vals = [i, comp, task, who, comm]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=F, size=10)
        cell.border = border
        cell.alignment = Alignment(horizontal="center" if c in (1, 4) else "left", vertical="center", wrap_text=True)
last = len(rows) + 1

for col, w in zip("ABCDE", (7, 26, 46, 24, 78)):
    ws.column_dimensions[col].width = w
ws.row_dimensions[1].height = 34
ws.freeze_panes = "A2"

dv = DataValidation(type="list", formula1='"Самостійно,AI,Спільно"', allow_blank=False)
dv.error = "Оберіть: Самостійно, AI або Спільно"
ws.add_data_validation(dv)
dv.add(f"D2:D{last}")

note_r = last + 2
ws.cell(row=note_r, column=1, value="Примітка: стовпець D має випадаючий список; аркуш «Підсумок» перераховує частки автоматично. Якщо фактичний внесок інший — змініть значення в стовпці D.").font = Font(name=F, italic=True, size=9)
ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=5)

# ---- Підсумок
sm = wb.create_sheet("Підсумок")
sm["A1"] = "Внесок у розробку (за кількістю задач)"
sm["A1"].font = Font(name=F, bold=True, size=12)
for c, h in enumerate(["Виконавець", "Кількість задач", "Частка"], 1):
    cell = sm.cell(row=3, column=c, value=h)
    cell.font = Font(name=F, bold=True)
    cell.fill = green
    cell.border = border
    cell.alignment = Alignment(horizontal="center")
for i, who in enumerate(["Самостійно", "AI", "Спільно"], 4):
    sm.cell(row=i, column=1, value=who)
    sm.cell(row=i, column=2, value=f"=COUNTIF('AI Interaction Log'!$D$2:$D${last},A{i})")
    sm.cell(row=i, column=3, value=f"=IF($B$7=0,0,B{i}/$B$7)")
    sm.cell(row=i, column=3).number_format = "0.0%"
sm["A7"] = "Разом"
sm["B7"] = "=SUM(B4:B6)"
sm["C7"] = "=SUM(C4:C6)"
sm["C7"].number_format = "0.0%"
for r in range(4, 8):
    for c in range(1, 4):
        cell = sm.cell(row=r, column=c)
        cell.font = Font(name=F, bold=(r == 7))
        cell.border = border
        if c > 1:
            cell.alignment = Alignment(horizontal="center")
sm["A9"] = "Рівень залучення AI: високий — AI виконав практично всю розробку; студент ставив завдання, надав варіант і матеріали, перевіряє й публікує результат."
sm["A9"].font = Font(name=F, italic=True, size=10)
sm.column_dimensions["A"].width = 22
sm.column_dimensions["B"].width = 18
sm.column_dimensions["C"].width = 12

from openpyxl.worksheet.properties import PageSetupProperties
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
wb.save("report/AI-Interaction-Log.xlsx")
print("saved", last)
